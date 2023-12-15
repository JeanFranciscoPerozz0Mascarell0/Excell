from openpyxl import Workbook
from openpyxl.chart import BarChart, series, Reference
from openpyxl.styles import Font

wb = Workbook()
ws = wb.active

salgados = [['Fritos', 'Sabor','Quantidade'], ['Pastel', 'Bom', 20], ['Risoles', 'Bom', 60], ['Enroladinho de salsicha', 'Bom', 120]]

for row in salgados:
    ws.append(row)
    ft = Font(bold=True)
    for row in ws['A1:B1']:
        for cell in row:
            cell.font = ft

chart = BarChart()
chart.title = "Salgados"
chart.type = "col"

data = Reference(ws, min_col=3, min_row=2, max_row=4, max_col=3)
categorias = Reference(ws, min_col=1, min_row=2, max_row=4, max_col=1)
chart.add_data(data)
chart.set_categories(categorias)

ws.add_chart(chart, 'd1')

wb.save('Graficos.xlsx')

